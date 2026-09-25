import os
import smtplib
import hashlib
from datetime import datetime
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

import streamlit as st
from sqlalchemy import create_engine, text

# -----------------------------------------------------------------------------
# CONFIGURACIÓN DE CONEXIÓN A BASE DE DATOS
# -----------------------------------------------------------------------------
def get_db_url():
    if "SUPABASE_DB_URL" in st.secrets:
        return st.secrets["SUPABASE_DB_URL"]
    return "sqlite:///dgcat_gestion.db"

def get_engine():
    return create_engine(get_db_url(), pool_pre_ping=True)

CORREO_DESTINO = "actmosaicocatastral@gmail.com"
SMTP_SERVER = "smtp.gmail.com"
SMTP_PORT = 587
SMTP_USER = os.getenv("SMTP_USER", "actmosaicocatastral@gmail.com")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "wycd ksbc mjbc onyw")

def hash_password(password: str) -> str:
    return hashlib.sha256(password.strip().encode()).hexdigest()

# -----------------------------------------------------------------------------
# AUDITORÍA DE OPERACIONES (INDEPENDIENTE)
# -----------------------------------------------------------------------------
def registrar_auditoria(oficio_id: int, usuario: str, accion: str, detalles: str = ""):
    """Registra una acción de auditoría sin borrar historial aunque se elimine el oficio."""
    try:
        engine = get_engine()
        with engine.begin() as conn:
            conn.execute(text("""
                INSERT INTO auditoria_oficios (oficio_id, usuario, accion, fecha_hora, detalles)
                VALUES (:of_id, :usr, :act, CURRENT_TIMESTAMP, :det)
            """), {
                "of_id": oficio_id,
                "usr": usuario or "SISTEMA",
                "act": accion,
                "det": detalles
            })
    except Exception as e:
        print(f"⚠️ Error al registrar auditoría: {e}")

# -----------------------------------------------------------------------------
# NOTIFICACIONES POR CORREO
# -----------------------------------------------------------------------------
def enviar_notificacion_correo(nombre_completo, username, rol):
    if not SMTP_PASSWORD:
        return
    try:
        msg = MIMEMultipart()
        msg['From'] = SMTP_USER
        msg['To'] = CORREO_DESTINO
        msg['Subject'] = f"🔔 Nuevo Registro de Usuario en Sistema DGCAT - {username.upper()}"

        rol_desc = "Administrador (Acceso Completo)" if rol == "admin" else ("Supervisor / Directivo" if rol == "supervisor" else "Capturista / Operador")
        fecha_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        cuerpo_html = f"""
        <html>
        <body style="font-family: Arial, sans-serif; color: #333333;">
            <div style="background-color: #064E3B; padding: 15px; text-align: center; color: white;">
                <h2>🏛️ DIRECCIÓN GENERAL DE CATASTRO (DGCAT)</h2>
                <p>Notificación Automática de Registro de Usuario</p>
            </div>
            <div style="border: 1px solid #10B981; padding: 20px; background-color: #F8FAFC;">
                <table style="width: 100%;">
                    <tr><td><b>Nombre Completo:</b></td><td>{nombre_completo}</td></tr>
                    <tr><td><b>Usuario:</b></td><td><code>{username}</code></td></tr>
                    <tr><td><b>Perfil:</b></td><td>{rol_desc}</td></tr>
                    <tr><td><b>Fecha:</b></td><td>{fecha_str}</td></tr>
                </table>
            </div>
        </body>
        </html>
        """
        msg.attach(MIMEText(cuerpo_html, 'html'))
        server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.sendmail(SMTP_USER, CORREO_DESTINO, msg.as_string())
        server.quit()
    except Exception as e:
        print(f"⚠️ Error en envío de correo: {e}")

# -----------------------------------------------------------------------------
# INICIALIZACIÓN DE BASE DE DATOS
# -----------------------------------------------------------------------------
def init_db():
    engine = get_engine()
    with engine.begin() as conn:
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS usuarios (
                id SERIAL PRIMARY KEY,
                username VARCHAR(100) UNIQUE NOT NULL,
                password_hash VARCHAR(255) NOT NULL,
                nombre_completo VARCHAR(255) NOT NULL,
                rol VARCHAR(50) DEFAULT 'operador',
                password_plain TEXT
            );
        """))

        # Usuarios iniciales
        res_admin = conn.execute(text("SELECT COUNT(*) FROM usuarios WHERE username = 'admin'")).scalar()
        if res_admin == 0:
            conn.execute(text("""
                INSERT INTO usuarios (username, password_hash, nombre_completo, rol, password_plain)
                VALUES ('admin', :p, 'Administrador DGCAT', 'admin', 'admin123')
            """), {"p": hash_password("admin123")})

        # Tablas de catálogos
        conn.execute(text("CREATE TABLE IF NOT EXISTS cat_scg (nombre VARCHAR(100) UNIQUE);"))
        conn.execute(text("CREATE TABLE IF NOT EXISTS cat_siscat (nombre VARCHAR(100) UNIQUE);"))
        conn.execute(text("CREATE TABLE IF NOT EXISTS cat_sistemas_or (nombre VARCHAR(100) UNIQUE);"))
        conn.execute(text("CREATE TABLE IF NOT EXISTS cat_tramite (nombre VARCHAR(100) UNIQUE);"))

        # Inserción inicial de catálogos
        for item in ["BANDEJA DE GEOGRAFO", "CON RESPUESTA PREVIA", "CONCLUIDO", "EN ESPERA DE SISTEMAS", "EN OTRA BANDEJA", "GEOG. PATRICIA", "SISTEMAS", "SUBIDO"]:
            conn.execute(text("INSERT INTO cat_scg (nombre) VALUES (:n) ON CONFLICT DO NOTHING;"), {"n": item})
        for item in ["ACUSE", "CONCLUIDO", "CORREO", "SISTEMAS", "SUBIDO"]:
            conn.execute(text("INSERT INTO cat_siscat (nombre) VALUES (:n) ON CONFLICT DO NOTHING;"), {"n": item})
        for item in ["SISTEMAS", "OR"]:
            conn.execute(text("INSERT INTO cat_sistemas_or (nombre) VALUES (:n) ON CONFLICT DO NOTHING;"), {"n": item})
        for item in ["CAMBIO DE DESTINO", "CAMBIO DE SUPERFICIE", "DOMINIO PLENO", "ACT. DE MOSAICO", "SENTENCIA"]:
            conn.execute(text("INSERT INTO cat_tramite (nombre) VALUES (:n) ON CONFLICT DO NOTHING;"), {"n": item})

        # Tabla principal de oficios
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS oficios (
                id SERIAL PRIMARY KEY,
                id_registro VARCHAR(50),
                estado VARCHAR(100),
                municipio VARCHAR(100),
                ejido VARCHAR(200),
                no_oficio VARCHAR(100),
                dgcat VARCHAR(100),
                fecha_entrega VARCHAR(50),
                fecha_recibido VARCHAR(50),
                scg VARCHAR(100),
                siscat VARCHAR(100),
                sistemas_or VARCHAR(100),
                tipo_tramite VARCHAR(100),
                observaciones TEXT,
                archivo_escaneado VARCHAR(255),
                creado_por VARCHAR(100),
                fecha_creacion TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP,
                modificado_por VARCHAR(100),
                fecha_modificacion TIMESTAMP WITH TIME ZONE
            );
        """))

        # Tabla de seguimiento de predio
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS seguimiento_predio (
                id SERIAL PRIMARY KEY,
                dgcat VARCHAR(100),
                estado VARCHAR(100),
                municipio VARCHAR(100),
                ejido VARCHAR(200),
                fecha_registro VARCHAR(50),
                fecha_actualizacion VARCHAR(50),
                observaciones TEXT,
                archivo_escaneado VARCHAR(255),
                registrado_por VARCHAR(100)
            );
        """))

        # Tabla de Auditoría
        conn.execute(text("""
            CREATE TABLE IF NOT EXISTS auditoria_oficios (
                id SERIAL PRIMARY KEY,
                oficio_id INT,
                usuario VARCHAR(100) NOT NULL,
                accion VARCHAR(50) NOT NULL,
                fecha_hora TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP,
                detalles TEXT
            );
        """))

# -----------------------------------------------------------------------------
# PAGINACIÓN Y BÚSQUEDA DEL LADO DEL BACKEND
# -----------------------------------------------------------------------------
def obtener_oficios_paginados(page: int = 1, page_size: int = 200, busqueda: str = "", estado: str = "TODOS", scg: str = "TODOS", siscat: str = "TODOS"):
    """Consulta paginada eficiente con LIMIT y OFFSET."""
    engine = get_engine()
    offset = (page - 1) * page_size
    
    where_clauses = ["1=1"]
    params = {"limit": page_size, "offset": offset}

    if busqueda and busqueda.strip():
        where_clauses.append("(UPPER(dgcat) LIKE :b OR UPPER(no_oficio) LIKE :b OR UPPER(estado) LIKE :b OR UPPER(ejido) LIKE :b)")
        params["b"] = f"%{busqueda.strip().upper()}%"
    if estado != "TODOS":
        where_clauses.append("UPPER(estado) = :edo")
        params["edo"] = estado.upper()
    if scg != "TODOS":
        where_clauses.append("UPPER(scg) = :scg")
        params["scg"] = scg.upper()
    if siscat != "TODOS":
        where_clauses.append("UPPER(siscat) = :siscat")
        params["siscat"] = siscat.upper()

    where_sql = " AND ".join(where_clauses)

    query_count = f"SELECT COUNT(*) FROM oficios WHERE {where_sql}"
    query_data = f"SELECT * FROM oficios WHERE {where_sql} ORDER BY id DESC LIMIT :limit OFFSET :offset"

    with engine.connect() as conn:
        total_records = conn.execute(text(query_count), params).scalar()
        df = pd.read_sql(text(query_data), conn, params=params)

    total_pages = max(1, (total_records + page_size - 1) // page_size)
    return df, total_records, total_pages

def obtener_seguimiento_paginado(page: int = 1, page_size: int = 200, busqueda: str = ""):
    engine = get_engine()
    offset = (page - 1) * page_size
    where_sql = "1=1"
    params = {"limit": page_size, "offset": offset}

    if busqueda and busqueda.strip():
        where_sql += " AND (UPPER(dgcat) LIKE :b OR UPPER(estado) LIKE :b OR UPPER(municipio) LIKE :b OR UPPER(ejido) LIKE :b)"
        params["b"] = f"%{busqueda.strip().upper()}%"

    query_count = f"SELECT COUNT(*) FROM seguimiento_predio WHERE {where_sql}"
    query_data = f"SELECT * FROM seguimiento_predio WHERE {where_sql} ORDER BY id DESC LIMIT :limit OFFSET :offset"

    with engine.connect() as conn:
        total_records = conn.execute(text(query_count), params).scalar()
        df = pd.read_sql(text(query_data), conn, params=params)

    total_pages = max(1, (total_records + page_size - 1) // page_size)
    return df, total_records, total_pages

# -----------------------------------------------------------------------------
# FUNCIONES DE MANTENIMIENTO Y AUDITORÍA
# -----------------------------------------------------------------------------
def existe_folio_oficio(dgcat_folio: str, excluir_id: int = None):
    engine = get_engine()
    with engine.connect() as conn:
        if excluir_id:
            res = conn.execute(text("SELECT id FROM oficios WHERE UPPER(TRIM(dgcat)) = :dg AND id != :id"), {"dg": dgcat_folio.strip().upper(), "id": excluir_id}).fetchone()
        else:
            res = conn.execute(text("SELECT id FROM oficios WHERE UPPER(TRIM(dgcat)) = :dg"), {"dg": dgcat_folio.strip().upper()}).fetchone()
        return res[0] if res else None

def existe_folio_seguimiento(dgcat_folio: str, excluir_id: int = None):
    engine = get_engine()
    with engine.connect() as conn:
        if excluir_id:
            res = conn.execute(text("SELECT id FROM seguimiento_predio WHERE UPPER(TRIM(dgcat)) = :dg AND id != :id"), {"dg": dgcat_folio.strip().upper(), "id": excluir_id}).fetchone()
        else:
            res = conn.execute(text("SELECT id FROM seguimiento_predio WHERE UPPER(TRIM(dgcat)) = :dg"), {"dg": dgcat_folio.strip().upper()}).fetchone()
        return res[0] if res else None

def guardar_oficio(datos: dict, id_oficio: int = None, usuario_actual: str = "SISTEMA"):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            if id_oficio:
                # Modificación
                query = text("""
                    UPDATE oficios
                    SET id_registro = :id_registro, estado = :estado, municipio = :municipio, ejido = :ejido,
                        no_oficio = :no_oficio, dgcat = :dgcat, fecha_entrega = :fecha_entrega, fecha_recibido = :fecha_recibido,
                        scg = :scg, siscat = :siscat, sistemas_or = :sistemas_or, tipo_tramite = :tipo_tramite,
                        observaciones = :observaciones, 
                        archivo_escaneado = COALESCE(NULLIF(:archivo_escaneado, ''), archivo_escaneado),
                        modificado_por = :usr, fecha_modificacion = CURRENT_TIMESTAMP
                    WHERE id = :id
                """)
                datos["usr"] = usuario_actual
                datos["id"] = id_oficio
                conn.execute(query, datos)
                
                # Registrar Auditoría Independiente
                registrar_auditoria(id_oficio, usuario_actual, "MODIFICACION", f"Oficio {datos.get('dgcat')} modificado.")
                return True, f"Oficio {datos.get('dgcat')} actualizado exitosamente."
            else:
                # Nuevo Registro
                query = text("""
                    INSERT INTO oficios (id_registro, estado, municipio, ejido, no_oficio, dgcat, fecha_entrega, fecha_recibido, scg, siscat, sistemas_or, tipo_tramite, observaciones, archivo_escaneado, creado_por, fecha_creacion)
                    VALUES (:id_registro, :estado, :municipio, :ejido, :no_oficio, :dgcat, :fecha_entrega, :fecha_recibido, :scg, :siscat, :sistemas_or, :tipo_tramite, :observaciones, :archivo_escaneado, :usr, CURRENT_TIMESTAMP)
                    RETURNING id
                """)
                datos["usr"] = usuario_actual
                new_id = conn.execute(query, datos).scalar()
                
                # Registrar Auditoría Independiente
                registrar_auditoria(new_id, usuario_actual, "CREACION", f"Nuevo oficio {datos.get('dgcat')} creado.")
                return True, f"Oficio {datos.get('dgcat')} guardado exitosamente."
    except Exception as e:
        return False, f"Error al guardar oficio: {e}"

def eliminar_oficio(oficio_id: int, usuario_actual: str = "SISTEMA"):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            # Obtener datos previos para auditoría
            res = conn.execute(text("SELECT dgcat FROM oficios WHERE id = :id"), {"id": oficio_id}).fetchone()
            folio_dgcat = res[0] if res else str(oficio_id)

            conn.execute(text("DELETE FROM oficios WHERE id = :id"), {"id": oficio_id})
            
            # Auditoría independiente que PERMANECE tras borrar
            registrar_auditoria(oficio_id, usuario_actual, "ELIMINACION", f"Oficio {folio_dgcat} (ID #{oficio_id}) eliminado permanentemente.")
            return True, f"El oficio ID #{oficio_id} ({folio_dgcat}) ha sido eliminado."
    except Exception as e:
        return False, f"Error al eliminar oficio: {e}"

def guardar_seguimiento_predio(datos: dict, id_registro: int = None):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            if id_registro:
                query = text("""
                    UPDATE seguimiento_predio
                    SET estado = :estado, municipio = :municipio, ejido = :ejido,
                        fecha_actualizacion = :fecha_actualizacion, observaciones = :observaciones,
                        archivo_escaneado = COALESCE(NULLIF(:archivo_escaneado, ''), archivo_escaneado),
                        registrado_por = :registrado_por
                    WHERE id = :id
                """)
                datos["id"] = id_registro
                conn.execute(query, datos)
                return True, "Seguimiento de predio actualizado correctamente."
            else:
                query = text("""
                    INSERT INTO seguimiento_predio (dgcat, estado, municipio, ejido, fecha_registro, fecha_actualizacion, observaciones, archivo_escaneado, registrado_por)
                    VALUES (:dgcat, :estado, :municipio, :ejido, :fecha_registro, :fecha_actualizacion, :observaciones, :archivo_escaneado, :registrado_por)
                """)
                conn.execute(query, datos)
                return True, "Seguimiento de predio registrado correctamente."
    except Exception as e:
        return False, f"Error al guardar seguimiento: {e}"

def eliminar_seguimiento_predio(id_registro: int):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            conn.execute(text("DELETE FROM seguimiento_predio WHERE id = :id"), {"id": id_registro})
            return True, f"Registro de seguimiento #{id_registro} eliminado."
    except Exception as e:
        return False, f"Error al eliminar seguimiento: {e}"

def verificar_login(username, password):
    engine = get_engine()
    pwd_hash = hash_password(password)
    with engine.connect() as conn:
        return conn.execute(
            text("SELECT username, nombre_completo, rol FROM usuarios WHERE LOWER(username) = LOWER(:u) AND password_hash = :p"),
            {"u": username.strip(), "p": pwd_hash}
        ).fetchone()

def registrar_nuevo_usuario(username, password, nombre_completo, rol="operador"):
    engine = get_engine()
    pwd_hash = hash_password(password)
    try:
        with engine.begin() as conn:
            conn.execute(
                text("INSERT INTO usuarios (username, password_hash, nombre_completo, rol, password_plain) VALUES (:u, :p, :n, :r, :pp)"),
                {"u": username.strip().lower(), "p": pwd_hash, "n": nombre_completo.strip(), "r": rol, "pp": password.strip()}
            )
        enviar_notificacion_correo(nombre_completo.strip(), username.strip().lower(), rol)
        return True, "Usuario registrado exitosamente."
    except Exception as e:
        return False, f"Error al registrar usuario: {e}"

def cambiar_password_usuario(username, nueva_password):
    engine = get_engine()
    pwd_hash = hash_password(nueva_password)
    try:
        with engine.begin() as conn:
            conn.execute(
                text("UPDATE usuarios SET password_hash = :p, password_plain = :pp WHERE LOWER(username) = LOWER(:u)"),
                {"p": pwd_hash, "pp": nueva_password.strip(), "u": username.strip()}
            )
        return True, f"Contraseña de '{username}' actualizada."
    except Exception as e:
        return False, f"Error al cambiar contraseña: {e}"

def eliminar_usuario(username):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            conn.execute(text("DELETE FROM usuarios WHERE LOWER(username) = LOWER(:u)"), {"u": username.strip()})
        return True, f"Usuario '{username}' eliminado correctamente."
    except Exception as e:
        return False, f"Error al eliminar usuario: {e}"

def actualizar_opcion_catalogo(tabla: str, valor_antiguo: str, valor_nuevo: str):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            conn.execute(text(f"UPDATE {tabla} SET nombre = :n WHERE nombre = :a"), {"n": valor_nuevo, "a": valor_antiguo})
        return True, f"Opción actualizada a '{valor_nuevo}'."
    except Exception as e:
        return False, f"Error al actualizar catálogo: {e}"

def contar_oficios_con_valor_catalogo(tabla: str, valor: str):
    engine = get_engine()
    col = "scg" if "scg" in tabla else ("siscat" if "siscat" in tabla else ("sistemas_or" if "sistemas_or" in tabla else "tipo_tramite"))
    with engine.connect() as conn:
        return conn.execute(text(f"SELECT COUNT(*) FROM oficios WHERE UPPER({col}) = :v"), {"v": valor.upper()}).scalar() or 0

def eliminar_opcion_catalogo(tabla: str, valor: str):
    engine = get_engine()
    try:
        with engine.begin() as conn:
            conn.execute(text(f"DELETE FROM {tabla} WHERE nombre = :v"), {"v": valor})
        return True, f"Opción '{valor}' eliminada del catálogo."
    except Exception as e:
        return False, f"Error al eliminar del catálogo: {e}"

# -----------------------------------------------------------------------------
# GENERADOR DE EXCEL EJECUTIVO (OCULTA ID Y CENTRA LA GRÁFICA)
# -----------------------------------------------------------------------------
def generar_excel_ejecutivo(df, filename="Reporte_DGCAT_Ejecutivo.xlsx"):
    wb = openpyxl.Workbook()
    
    # 1. Resumen Ejecutivo
    ws_sum = wb.active
    ws_sum.title = "Resumen Ejecutivo"
    ws_sum.views.sheetView[0].showGridLines = True
    
    DARK_GREEN = "064E3B"
    HEADER_FILL = "0F172A"
    BORDER_COLOR = "CBD5E1"
    
    font_title = Font(name="Calibri", size=15, bold=True, color=DARK_GREEN)
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_kpi_num = Font(name="Calibri", size=20, bold=True, color=DARK_GREEN)
    font_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="475569")
    
    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR), right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR), bottom=Side(style='thin', color=BORDER_COLOR)
    )
    
    ws_sum["A1"] = "DIRECCIÓN GENERAL DE CATASTRO - REGISTRO AGRARIO NACIONAL"
    ws_sum["A1"].font = font_title
    ws_sum["A2"] = f"Reporte Ejecutivo | Consulta realizada el {datetime.now().strftime('%d de %B de %Y')}"
    ws_sum["A2"].font = font_subtitle
    
    total = len(df)
    scg_concluidos = len(df[df['scg'] == 'CONCLUIDO']) if 'scg' in df.columns else 0
    siscat_subidos = len(df[df['siscat'].isin(['CONCLUIDO', 'SUBIDO'])]) if 'siscat' in df.columns else 0
    
    kpis = [("TOTAL OFICIOS", total), ("CONCLUIDOS SCG", scg_concluidos), ("PROCESADOS SISCAT", siscat_subidos)]
    
    col_idx = 1
    for label, val in kpis:
        cell_lbl = ws_sum.cell(row=4, column=col_idx, value=label)
        cell_lbl.font = font_kpi_lbl
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        
        cell_val = ws_sum.cell(row=5, column=col_idx, value=val)
        cell_val.font = font_kpi_num
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
        
        ws_sum.merge_cells(start_row=4, start_column=col_idx, end_row=4, end_column=col_idx+1)
        ws_sum.merge_cells(start_row=5, start_column=col_idx, end_row=5, end_column=col_idx+1)
        
        for r in range(4, 6):
            for c in range(col_idx, col_idx+2):
                ws_sum.cell(row=r, column=c).border = thin_border
        col_idx += 3
        
    ws_sum.cell(row=8, column=1, value="Estatus por Bandeja SCG").font = Font(name="Calibri", size=12, bold=True, color=DARK_GREEN)
    ws_sum.cell(row=9, column=1, value="Bandeja").font = font_header
    ws_sum.cell(row=9, column=1).fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    ws_sum.cell(row=9, column=2, value="Cantidad").font = font_header
    ws_sum.cell(row=9, column=2).fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    
    r_idx = 10
    start_chart_row = r_idx
    if 'scg' in df.columns and not df.empty:
        scg_counts = df['scg'].value_counts()
        for b_name, b_cnt in scg_counts.items():
            c1 = ws_sum.cell(row=r_idx, column=1, value=str(b_name))
            c2 = ws_sum.cell(row=r_idx, column=2, value=int(b_cnt))
            c1.border = thin_border
            c2.border = thin_border
            if r_idx % 2 == 0:
                c1.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                c2.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            r_idx += 1

    # Insertar Gráfica Centrada (Ajuste Punto 4)
    if r_idx > start_chart_row:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Volumetría por Bandeja SCG"
        chart.y_axis.title = "Cantidad"
        chart.x_axis.title = "Bandeja"
        
        data_ref = Reference(ws_sum, min_col=2, min_row=9, max_row=r_idx-1)
        cats_ref = Reference(ws_sum, min_col=1, min_row=10, max_row=r_idx-1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 16
        chart.height = 10
        
        # Posicionar centrada al lado de la tabla
        ws_sum.add_chart(chart, "D8")

    # 2. Detalle de Oficios (OCULTA EL ID INTERNO)
    ws_det = wb.create_sheet(title="Detalle de Oficios")
    ws_det.views.sheetView[0].showGridLines = True
    
    header_map = {
        'id_registro': 'ID Numérico',
        'estado': 'Estado / Entidad',
        'municipio': 'Municipio',
        'ejido': 'Ejido / Núcleo Agrario',
        'no_oficio': 'No. de Oficio',
        'dgcat': 'Folio DGCAT',
        'fecha_entrega': 'Fecha Entrega',
        'fecha_recibido': 'Fecha Recibido',
        'scg': 'Bandeja SCG',
        'siscat': 'Estatus SISCAT',
        'sistemas_or': 'SISTEMAS/OR',
        'tipo_tramite': 'Tipo de Trámite',
        'observaciones': 'Observaciones',
        'archivo_escaneado': 'Expediente PDF'
    }
    cols = [c for c in df.columns if c in header_map]
    
    ws_det.cell(row=1, column=1, value="REGISTRO DETALLADO DE OFICIOS DE RESPUESTA - DGCAT").font = font_title
    
    for c_idx, col_name in enumerate(cols, start=1):
        cell = ws_det.cell(row=3, column=c_idx, value=header_map[col_name])
        cell.font = font_header
        cell.fill = PatternFill(start_color=DARK_GREEN, end_color=DARK_GREEN, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    ws_det.row_dimensions[3].height = 26
    
    for r_offset, row in df.iterrows():
        row_num = r_offset + 4
        is_zebra = (row_num % 2 == 0)
        zebra_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid") if is_zebra else PatternFill(fill_type=None)
        
        for c_idx, col_name in enumerate(cols, start=1):
            val = row[col_name]
            val_str = "" if pd.isna(val) else str(val).strip()
            cell = ws_det.cell(row=row_num, column=c_idx, value=val_str)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill
            cell.alignment = Alignment(horizontal="center" if col_name in ['id_registro', 'fecha_entrega', 'fecha_recibido', 'scg', 'siscat'] else "left", vertical="center")

    for ws in [ws_sum, ws_det]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    wb.save(filename)
    return filename

def generar_excel_seguimiento(df, filename="Reporte_Seguimiento_Predio.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Seguimiento de Predio"
    ws.views.sheetView[0].showGridLines = True

    header_map = {
        'dgcat': 'Folio DGCAT',
        'estado': 'Estado',
        'municipio': 'Municipio',
        'ejido': 'Ejido / Núcleo Agrario',
        'fecha_registro': 'Fecha Registro',
        'fecha_actualizacion': 'Última Actualización',
        'observaciones': 'Observaciones',
        'archivo_escaneado': 'Expediente PDF',
        'registrado_por': 'Registrado Por'
    }
    cols = [c for c in df.columns if c in header_map]

    ws.cell(row=1, column=1, value=f"REPORTE DE SEGUIMIENTO DE PREDIO | Consulta: {datetime.now().strftime('%d/%m/%Y')}").font = Font(size=14, bold=True, color="064E3B")

    for c_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=c_idx, value=header_map[col_name])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")

    for r_offset, row in df.iterrows():
        row_num = r_offset + 4
        for c_idx, col_name in enumerate(cols, start=1):
            val = row[col_name]
            ws.cell(row=row_num, column=c_idx, value="" if pd.isna(val) else str(val).strip())

    wb.save(filename)
    return filename
